"""Self-check for the import wizard's file parsing and column mapping:
`python test_contact_import_wizard.py` (no pytest, no network, no database).

Builds a workbook shaped like the real one — JOB AGENCIES UK (Company, Link
Job/Company, LinkedIn, Phone, Email; company groups separated by blank rows)
and UK COMPANIES (COMPANY, JOB TITLE/LINK) — with real cell hyperlinks, and
checks what reaches the database: raw cells preserved, blank-row grouping,
LinkedIn display text never taken as a URL, no guessed company names, and job
text kept as source information. Matching itself is tested in SQL
(supabase/tests/048_contact_import.test.sql).
"""

from __future__ import annotations

import io

import openpyxl

from app.contacts_import.mapping import apply_mapping, valid_url, validate_mapping
from app.contacts_import.upload import MAX_FILE_BYTES, ROLES, UploadError, parse_upload
from app.routers.contact_imports import with_blank_rows

_failures: list[str] = []


def check(name: str, cond: bool, detail: object = "") -> None:
    print(f"[{'ok  ' if cond else 'FAIL'}] {name}{('  ' + repr(detail)) if detail != '' and not cond else ''}")
    if not cond:
        _failures.append(name)


def _set(ws, ref: str, value, link: str | None = None) -> None:
    ws[ref] = value
    if link:
        ws[ref].hyperlink = link


def build_workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "JOB AGENCIES UK"
    for col, header in zip("ABCDE", ["Company", "Link Job/Company", "LinkedIn", "Phone", "Email"]):
        ws[f"{col}1"] = header
    _set(ws, "A2", "Acme Recruitment", "https://www.acme-rec.co.uk")
    _set(ws, "B2", "Recruiter role", "https://www.linkedin.com/jobs/view/4418807838/")
    _set(ws, "C2", "Jane Doe - Senior Recruiter | LinkedIn", "https://www.linkedin.com/in/jane-doe/")
    ws["D2"] = 1753621902.0            # Excel ate the leading 0 of "01753 621902"
    ws["E2"] = "jane@acme-rec.co.uk"
    ws["B3"] = "no suitable roles available"
    _set(ws, "C3", "Tom Bell", "https://uk.linkedin.com/in/tombell")
    ws["D3"] = "07700 900123 (general)"
    ws["C4"] = "https://www.linkedin.com/in/someone"   # display text is a URL; no name
    ws["E4"] = "x@y.co"
    # row 5 completely blank: ends the Acme group
    _set(ws, "C6", "Orphan Person", "https://www.linkedin.com/in/orphan")
    ws["D6"] = "0123456789"
    ws["A7"] = "Beta Staffing"
    ws["C7"] = "Ann One\nBob Two"
    ws["D7"] = "0200000000"
    ws["C8"] = "linkedin.com/in/jane"                  # looks like a link, isn't a valid one, isn't a name
    ws["E8"] = "jane2@beta.com"
    _set(ws, "C9", "Carl Day", "https://www.linkedin.com/company/beta")
    ws["E9"] = "carl@beta"
    _set(ws, "A10", "https://www.get-recruited.co.uk/", "https://www.get-recruited.co.uk/")
    _set(ws, "C10", "Dee Ray", "https://www.linkedin.com/in/deeray")
    ws["E10"] = "dee@get-recruited.co.uk"

    uk = wb.create_sheet("UK COMPANIES")
    uk["A1"], uk["B1"] = "COMPANY", "JOB TITLE/LINK"
    uk["A2"] = "Gamma Ltd"
    _set(uk, "B2", "Software Engineer", "https://uk.indeed.com/viewjob?jk=abc")
    uk["A3"] = "Delta plc"
    # row 4 blank
    uk["B5"] = "Orphan Job"
    uk["A6"] = "Epsilon"
    uk["B6"] = "https://careers.epsilon.example/job/1"
    uk["A7"] = "Zeta"
    uk["B7"] = "Data Analyst"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def stored_rows(sheet) -> list[dict]:
    """What the upload route stores: non-blank rows only, under their sheet row numbers."""
    return [{"row_number": r.row_number, "raw": r.cells} for r in sheet.rows if not r.is_blank]


def by_number(results: list[dict]) -> dict[int, dict]:
    return {r["row_number"]: r for r in results}


kind, sheets = parse_upload("AI TEAM - CURRENT CONTACTS.xlsx", build_workbook())
agency, uk = sheets

# ── Parsing ──────────────────────────────────────────────────────────────────
check("xlsx: both sheets, in workbook order", kind == "xlsx" and [s.name for s in sheets] == ["JOB AGENCIES UK", "UK COMPANIES"])
check("headers read from row 1", agency.headers == {"A": "Company", "B": "Link Job/Company", "C": "LinkedIn", "D": "Phone", "E": "Email"})
row2 = agency.rows[0].cells
check("raw: a numeric phone stays a number, as stored in the file (openpyxl reads whole numbers as int)",
      row2["D"]["type"] == "int" and row2["D"]["value"] == 1753621902, row2["D"])
check("raw: the cell hyperlink is kept", row2["C"]["hyperlink"] == "https://www.linkedin.com/in/jane-doe/")
check("raw: display text is kept verbatim", row2["C"]["text"] == "Jane Doe - Senior Recruiter | LinkedIn")
check("blank rows counted, not stored", agency.summary()["blank_rows"] == 1 and agency.summary()["row_count"] == 8)

# ── Suggested mapping ────────────────────────────────────────────────────────
am, um = agency.suggested_mapping(), uk.suggested_mapping()
check("agency sheet: columns suggested from headers",
      am["columns"] == {"A": "company_name", "B": "source", "C": "contact", "D": "phone", "E": "email"}, am["columns"])
check("agency sheet: grouping and company-link-as-website suggested", am["group_fill"] and am["company_link_is_website"])
check("UK COMPANIES: company + source; no grouping, company link not a website",
      um["columns"] == {"A": "company_name", "B": "source"} and not um["group_fill"] and not um["company_link_is_website"], um)

# ── Mapping the agency sheet ─────────────────────────────────────────────────
a = by_number(apply_mapping(with_blank_rows(stored_rows(agency)), am))
check("every non-blank row produces a record; the blank row does not", sorted(a) == [2, 3, 4, 6, 7, 8, 9, 10], sorted(a))

r2 = a[2]["extracted"]
check("row 2: company name, slug and its own site", r2["company"] == {
    "name": "Acme Recruitment", "slug": "acme-recruitment", "website": "https://www.acme-rec.co.uk", "group_row": None}, r2["company"])
check("row 2: name and title split; ' | LinkedIn' wrapper dropped",
      r2["contact"]["full_name"] == "Jane Doe" and r2["contact"]["job_title"] == "Senior Recruiter", r2["contact"])
check("row 2: the hyperlink (a valid profile URL) becomes the LinkedIn URL",
      r2["contact"]["linkedin_url"] == "https://www.linkedin.com/in/jane-doe/")
check("row 2: Excel float phone gets its leading 0 back", r2["contact"]["phone"] == "01753621902", r2["contact"]["phone"])
check("row 2: the job link is source information, not a job or a contact field",
      r2["source"] == {"text": "Recruiter role", "url": "https://www.linkedin.com/jobs/view/4418807838/"}, r2["source"])

r3 = a[3]["extracted"]
check("row 3: blank company cell continues the group above", r3["company"]["name"] == "Acme Recruitment" and r3["company"]["group_row"] == 2)
check("row 3: '(general)' cleaned from the phone and kept as a note",
      r3["contact"]["phone"] == "07700 900123" and "general" in (r3["contact"]["notes"] or ""), r3["contact"])
check("row 3: link text without a link is source text, not a contact note",
      r3["source"] == {"text": "no suitable roles available", "url": None} and "suitable" not in (r3["contact"]["notes"] or ""))

check("row 4: a LinkedIn URL as display text is not a name → contact error, company kept",
      a[4]["parse_errors"] == [] and a[4]["extracted"]["contact"] is None
      and "no readable name" in a[4]["extracted"]["contact_error"], a[4])
check("row 6: after a completely blank row, a blank company is an error — never guessed",
      a[6]["parse_errors"] == ["No company for this row: the company group ended at a blank row above it."]
      and a[6]["extracted"]["company"]["name"] is None, a[6])
check("row 7: a new company starts a new group", a[7]["extracted"]["company"]["name"] == "Beta Staffing")
check("row 7: two people in one cell → contact error (the shared phone can't be attributed)",
      "Multiple distinct contacts" in (a[7]["extracted"].get("contact_error") or ""), a[7]["extracted"])
check("row 8: link-like display text that isn't a valid URL is neither a URL nor a name",
      "link-like text" in (a[8]["extracted"].get("contact_error") or "") and a[8]["extracted"]["company"]["name"] == "Beta Staffing", a[8])
check("row 9: an invalid email and a LinkedIn company page leave no contact method → contact error",
      "No usable email, phone or LinkedIn profile" in (a[9]["extracted"].get("contact_error") or ""), a[9]["extracted"])
r10 = a[10]["extracted"]
check("row 10: a URL-only company cell gives a website, never a guessed name",
      r10["company"]["name"] is None and r10["company"]["slug"] is None
      and r10["company"]["website"] == "https://www.get-recruited.co.uk/" and a[10]["parse_errors"] == [], r10["company"])
check("row 10: its contact is still read", r10["contact"]["full_name"] == "Dee Ray" and r10["contact"]["email"] == "dee@get-recruited.co.uk")

# ── Mapping UK COMPANIES ─────────────────────────────────────────────────────
u = by_number(apply_mapping(with_blank_rows(stored_rows(uk)), um))
check("UK row 2: company only; a job-board link is not its website",
      u[2]["extracted"]["company"]["name"] == "Gamma Ltd" and u[2]["extracted"]["company"]["website"] is None
      and u[2]["extracted"]["contact"] is None)
check("UK row 2: job title + link kept as source", u[2]["extracted"]["source"] == {"text": "Software Engineer", "url": "https://uk.indeed.com/viewjob?jk=abc"})
check("UK row 3: a company name alone is a valid row", u[3]["parse_errors"] == [] and u[3]["extracted"]["source"] is None)
check("UK row 5: no grouping on this sheet, so a missing company is an error", u[5]["parse_errors"] == ["No company name in this row."])
check("UK row 6: a real URL typed as text is the source URL", u[6]["extracted"]["source"] == {"text": None, "url": "https://careers.epsilon.example/job/1"})
check("UK row 7: plain job text stays source text", u[7]["extracted"]["source"] == {"text": "Data Analyst", "url": None})

excluded = apply_mapping(with_blank_rows(stored_rows(uk)), {**um, "include": False})
check("an excluded sheet's rows are marked excluded", all(r["extracted"] == {"excluded": True} for r in excluded) and len(excluded) == 5)

# ── Grouping edge cases ──────────────────────────────────────────────────────
rows = [
    {"row_number": 2, "cells": {"C": {"text": "Early Bird", "hyperlink": "https://www.linkedin.com/in/early"}}},
    {"row_number": 3, "cells": {"A": {"text": "Kappa"}, "C": {"text": "Kim", "hyperlink": "https://www.linkedin.com/in/kim"}}},
    {"row_number": 4, "cells": {"C": {"text": "Kit", "hyperlink": "https://www.linkedin.com/in/kit"}}},
    {"row_number": 5, "cells": {"C": {"text": "Kay", "hyperlink": "https://www.linkedin.com/in/kay"}}},
]
g = by_number(apply_mapping(rows, {"group_fill": True, "columns": {"A": "company_name", "C": "contact"}}))
check("grouping: a row before any company is an error", g[2]["parse_errors"] == ["No company for this row and none above it to continue."])
check("grouping: every row until the next blank row inherits", g[4]["extracted"]["company"]["name"] == "Kappa" and g[5]["extracted"]["company"]["group_row"] == 3)
gapped = with_blank_rows([{"row_number": 3, "raw": {"A": {"text": "X"}}}, {"row_number": 6, "raw": {"C": {"text": "Y"}}}])
check("a gap in stored row numbers is restored as one blank row", [r["row_number"] for r in gapped] == [3, 4, 6] and gapped[1]["cells"] == {})

# ── CSV ──────────────────────────────────────────────────────────────────────
csv_bytes = "Company;Website;Name;Email\nCafé Talent;cafe-talent.co.uk;Jo Lee;jo@cafe-talent.co.uk\n".encode("cp1252")
ckind, csheets = parse_upload("agencies.csv", csv_bytes)
cm = csheets[0].suggested_mapping()
c = apply_mapping(with_blank_rows(stored_rows(csheets[0])), cm)[0]["extracted"]
check("csv: cp1252 + semicolons parsed, mapping suggested",
      ckind == "csv" and cm["columns"] == {"A": "company_name", "B": "company_website", "C": "contact_name", "D": "email"}, cm["columns"])
check("csv: a bare domain in the website column becomes https://…; the name survives decoding",
      c["company"]["name"] == "Café Talent" and c["company"]["website"] == "https://cafe-talent.co.uk", c["company"])

# ── Upload limits and validation ─────────────────────────────────────────────
for label, name, data, needle in [
    ("wrong file type", "contacts.txt", b"x", "Upload an .xlsx"),
    ("empty file", "contacts.csv", b"", "empty"),
    ("over the size limit", "contacts.csv", b"x" * (MAX_FILE_BYTES + 1), "larger than"),
    ("not really an xlsx", "contacts.xlsx", b"PK not a zip", "not a valid .xlsx"),
]:
    try:
        parse_upload(name, data)
        check(f"upload refuses: {label}", False, "accepted")
    except UploadError as exc:
        check(f"upload refuses: {label}", needle in str(exc), str(exc))

# ── Formatting is not data; data is never dropped ────────────────────────────
from openpyxl.styles import PatternFill  # noqa: E402

from app.contacts_import import upload as upload_module  # noqa: E402


def small_book(fill=None) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append(["Company", "Email"])
    ws.append(["Acme", "a@acme.example"])
    ws.append(["Beta", "b@beta.example"])
    if fill:
        fill(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def refused(data: bytes, name: str = "f.xlsx") -> str:
    try:
        parse_upload(name, data)
    except UploadError as exc:
        return str(exc)
    return ""


def formatted_far_down(ws) -> None:
    ws["A30000"].fill = PatternFill("solid", fgColor="FFFF00")   # a fill dragged down, no value


_, fsheets = parse_upload("f.xlsx", small_book(formatted_far_down))
check("a formatted but empty cell at row 30,000 is not refused as 30,000 rows",
      [r.row_number for r in fsheets[0].rows] == [2, 3] and fsheets[0].summary()["blank_rows"] == 0,
      [r.row_number for r in fsheets[0].rows][-3:])
_, gsheets = parse_upload("f.xlsx", small_book(lambda ws: ws.__setitem__("A25000", "Gamma")))
check("real data far down keeps its row number, the gap kept as the blank rows it is",
      [r.row_number for r in gsheets[0].rows if not r.is_blank] == [2, 3, 25000])
saved_limit = upload_module.MAX_ROWS_PER_SHEET
upload_module.MAX_ROWS_PER_SHEET = 1
try:
    check("more data rows than the limit is still refused (formatted rows don't count, data rows do)",
          "more than 1 rows" in refused(small_book(formatted_far_down)), refused(small_book(formatted_far_down)))
finally:
    upload_module.MAX_ROWS_PER_SHEET = saved_limit

too_wide = refused(small_book(lambda ws: ws.__setitem__("AY3", "extra notes")))
check("data past column 50 is refused with the column named, not silently dropped",
      "column AY (row 3)" in too_wide and "first 50 columns" in too_wide, too_wide)
check("…a link alone past column 50 counts as data too",
      "column BA" in refused(small_book(lambda ws: setattr(ws["BA2"], "hyperlink", "https://www.linkedin.com/in/x"))))
_, wsheets = parse_upload("f.xlsx", small_book(lambda ws: setattr(ws["AY3"], "fill", PatternFill("solid", fgColor="FFFF00"))))
check("…but a merely formatted cell past column 50 is fine", len(wsheets[0].rows) == 2)
wide_csv = ("Company,Email" + "," * 49 + "\nAcme,a@acme.example" + "," * 48 + ",overflow\n").encode()
check("csv: a value past column 50 is refused", "column AY (row 2)" in refused(wide_csv, "w.csv"), refused(wide_csv, "w.csv"))
blank_tail_csv = ("Company,Email\nAcme,a@acme.example\n" + ",\n" * 30).encode()
upload_module.MAX_ROWS_PER_SHEET = 1
try:
    check("csv: rows of empty separators don't count against the row limit", refused(blank_tail_csv, "t.csv") == "")
finally:
    upload_module.MAX_ROWS_PER_SHEET = saved_limit

sheet_meta = [agency.summary(), uk.summary()]
check("mapping validation: accepts the suggestion",
      validate_mapping({"JOB AGENCIES UK": am, "UK COMPANIES": um}, sheet_meta, ROLES) is None)
for label, mapping, needle in [
    ("unknown sheet", {"Nope": am}, "Unknown sheet"),
    ("unknown role", {"UK COMPANIES": {**um, "columns": {"A": "company_name", "B": "salary"}}}, "Unknown role"),
    ("unknown column", {"UK COMPANIES": {**um, "columns": {"A": "company_name", "Z": "email"}}}, "no column"),
    ("no company column", {"UK COMPANIES": {**um, "columns": {"B": "source"}}}, "needs a company"),
    ("two company columns", {"UK COMPANIES": {**um, "columns": {"A": "company_name", "B": "company_name"}}}, "more than one"),
    ("nothing included", {"UK COMPANIES": {**um, "include": False}}, "at least one sheet"),
]:
    problem = validate_mapping(mapping, sheet_meta, ROLES)
    check(f"mapping validation refuses: {label}", bool(problem) and needle in problem, problem)

check("valid_url: display text must be a full http(s) URL",
      valid_url("https://www.linkedin.com/in/x") and not valid_url("linkedin.com/in/x") and not valid_url("Jane Doe | LinkedIn")
      and not valid_url("https://bad host.com"))

print()
if _failures:
    print(f"{len(_failures)} FAILED")
    raise SystemExit(1)
print("ALL CONTACT IMPORT WIZARD CHECKS PASSED")
