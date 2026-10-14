"""Reads an uploaded .xlsx or .csv into raw rows for the import wizard.

Nothing here interprets a value. Every non-empty cell is kept as it was read
— its value (with its Python type), its display text and its hyperlink — so
the database holds the source exactly (contact_import_rows.raw) and the
mapping step (mapping.py) can be re-run with a different mapping without the
file. The only decisions made here are structural: which row is the header,
which rows are completely blank, and a *suggested* column mapping the admin
confirms or changes.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import zipfile
from dataclasses import dataclass, field
from typing import Any, Optional

import openpyxl
from openpyxl.utils import get_column_letter

MAX_FILE_BYTES = 10 * 1024 * 1024
# An .xlsx is a zip: a small upload can inflate to gigabytes. Refuse before
# openpyxl ever sees it.
MAX_XLSX_UNCOMPRESSED_BYTES = 150 * 1024 * 1024
MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 20_000
MAX_COLUMNS = 50
HEADER_SEARCH_ROWS = 10
SAMPLE_ROWS = 5

ROLES = (
    "company_name",          # the company's name (a URL-only cell is not a name)
    "company_website",       # the company's own site
    "contact",               # a person's name, optionally hyperlinked to their LinkedIn profile
    "contact_name",          # a plain name column
    "contact_linkedin_url",  # a column of LinkedIn profile URLs
    "contact_job_title",     # the person's role
    "phone",
    "email",
    "source",                # job title / link / free text: kept as source information
    "ignore",
)


class UploadError(ValueError):
    """The file can't be read as an import (shown to the admin as a 422)."""


@dataclass
class RawRow:
    row_number: int                       # 1-based, as the spreadsheet shows it
    cells: dict[str, dict[str, Any]]      # column letter -> {header, value, text, type, hyperlink}

    @property
    def is_blank(self) -> bool:
        return not self.cells


@dataclass
class ParsedSheet:
    name: str
    index: int
    header_row: int
    headers: dict[str, str]               # column letter -> header text
    rows: list[RawRow] = field(default_factory=list)   # data rows, blank ones included

    def suggested_mapping(self) -> dict[str, Any]:
        return suggest_mapping(self)

    def summary(self) -> dict[str, Any]:
        data = [r for r in self.rows if not r.is_blank]
        return {
            "name": self.name,
            "index": self.index,
            "header_row": self.header_row,
            "headers": self.headers,
            "row_count": len(data),
            "blank_rows": len(self.rows) - len(data),
            "sample": [{"row_number": r.row_number, "cells": r.cells} for r in data[:SAMPLE_ROWS]],
            "suggested_mapping": self.suggested_mapping(),
        }


def _json_value(value: Any) -> tuple[Any, str]:
    """A JSON-safe copy of a cell value plus the name of its original type, so
    "01753 621902" typed as text and 1753621902 stored as a number stay
    distinguishable in the raw record."""
    if isinstance(value, bool):
        return value, "bool"
    if isinstance(value, int):
        return value, "int"
    if isinstance(value, float):
        return value, "float"
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat(), type(value).__name__
    return str(value), "str"


def _cell_record(value: Any, hyperlink: Optional[str], header: str) -> Optional[dict[str, Any]]:
    if value is None or (isinstance(value, str) and not value.strip()):
        if not hyperlink:
            return None
        value = ""
    json_value, kind = _json_value(value)
    text = json_value if isinstance(json_value, str) else str(json_value)
    return {"header": header, "value": json_value, "text": text, "type": kind, "hyperlink": hyperlink or None}


def _find_header(grid: list[list[tuple[Any, Optional[str]]]]) -> int:
    """Index of the first row with at least one non-empty cell, within the
    first few rows. Sheets in the wild sometimes start with a blank row."""
    for i, row in enumerate(grid[:HEADER_SEARCH_ROWS]):
        if any(v is not None and str(v).strip() for v, _ in row):
            return i
    raise UploadError("No header row found in the first rows of the sheet.")


def _sheet_from_grid(name: str, index: int, grid: list[list[tuple[Any, Optional[str]]]]) -> ParsedSheet:
    if not grid:
        return ParsedSheet(name=name, index=index, header_row=1, headers={})
    h = _find_header(grid)
    width = min(max(len(r) for r in grid), MAX_COLUMNS)
    headers: dict[str, str] = {}
    for c in range(width):
        value = grid[h][c][0] if c < len(grid[h]) else None
        text = re.sub(r"\s+", " ", str(value)).strip() if value is not None else ""
        headers[get_column_letter(c + 1)] = text or f"Column {get_column_letter(c + 1)}"

    rows: list[RawRow] = []
    for i in range(h + 1, len(grid)):
        cells: dict[str, dict[str, Any]] = {}
        for c in range(min(len(grid[i]), width)):
            letter = get_column_letter(c + 1)
            record = _cell_record(grid[i][c][0], grid[i][c][1], headers[letter])
            if record is not None:
                cells[letter] = record
        rows.append(RawRow(row_number=i + 1, cells=cells))

    # Trailing blank rows carry nothing (not even a group boundary worth keeping).
    while rows and rows[-1].is_blank:
        rows.pop()
    if len(rows) > MAX_ROWS_PER_SHEET:
        raise UploadError(f"Sheet {name!r} has more than {MAX_ROWS_PER_SHEET} rows. Split the file and import it in parts.")
    return ParsedSheet(name=name, index=index, header_row=h + 1, headers=headers, rows=rows)


def parse_xlsx(data: bytes) -> list[ParsedSheet]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            if sum(info.file_size for info in zf.infolist()) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise UploadError("The workbook is too large once uncompressed.")
    except zipfile.BadZipFile as exc:
        raise UploadError("This is not a valid .xlsx file.") from exc

    try:
        # Not read_only: read-only worksheets don't load hyperlinks, and the
        # hyperlink is often the only place a LinkedIn profile URL lives.
        # data_only: cached formula results, never formula text.
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as exc:  # openpyxl raises a zoo of exception types for corrupt files
        raise UploadError("The workbook could not be read.") from exc

    if len(wb.sheetnames) > MAX_SHEETS:
        raise UploadError(f"The workbook has more than {MAX_SHEETS} sheets.")

    sheets: list[ParsedSheet] = []
    for index, ws in enumerate(wb.worksheets):
        if ws.max_row > MAX_ROWS_PER_SHEET + HEADER_SEARCH_ROWS + 1000:
            raise UploadError(f"Sheet {ws.title!r} has more than {MAX_ROWS_PER_SHEET} rows. Split the file and import it in parts.")
        grid = [
            [(cell.value, cell.hyperlink.target if cell.hyperlink else None) for cell in row[:MAX_COLUMNS]]
            for row in ws.iter_rows()
        ]
        sheets.append(_sheet_from_grid(ws.title, index, grid))
    return sheets


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UploadError("The CSV file is not UTF-8 or Windows-1252 text.")


def parse_csv(data: bytes, sheet_name: str = "CSV") -> list[ParsedSheet]:
    text = _decode(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    grid = [[(value if value != "" else None, None) for value in row] for row in reader]
    if len(grid) > MAX_ROWS_PER_SHEET + HEADER_SEARCH_ROWS:
        raise UploadError(f"The CSV has more than {MAX_ROWS_PER_SHEET} rows. Split the file and import it in parts.")
    return [_sheet_from_grid(sheet_name, 0, grid)]


def parse_upload(filename: str, data: bytes) -> tuple[str, list[ParsedSheet]]:
    """Returns (file_kind, sheets)."""
    if len(data) > MAX_FILE_BYTES:
        raise UploadError(f"The file is larger than {MAX_FILE_BYTES // (1024 * 1024)} MB.")
    if not data:
        raise UploadError("The file is empty.")
    lower = filename.lower()
    if lower.endswith(".xlsx"):
        return "xlsx", parse_xlsx(data)
    if lower.endswith(".csv"):
        return "csv", parse_csv(data, sheet_name=re.sub(r"\.csv$", "", filename, flags=re.I) or "CSV")
    raise UploadError("Upload an .xlsx workbook or a .csv file.")


# ── Suggested mapping ─────────────────────────────────────────────────────────

_URL_RE = re.compile(r"^https?://", re.I)


def _norm_header(header: str) -> str:
    return re.sub(r"[^a-z]+", " ", header.lower()).strip()


def _column_samples(sheet: ParsedSheet, letter: str, limit: int = 20) -> list[dict[str, Any]]:
    out = []
    for row in sheet.rows:
        cell = row.cells.get(letter)
        if cell:
            out.append(cell)
            if len(out) >= limit:
                break
    return out


def _role_for(sheet: ParsedSheet, letter: str, header: str) -> str:
    h = _norm_header(header)
    words = set(h.split())
    if not h:
        return "ignore"
    if "link" in words or "links" in words:
        # "Link Job/Company", "JOB TITLE/LINK": a posting or page — source information.
        return "source"
    if "linkedin" in words:
        samples = _column_samples(sheet, letter)
        # Display text that is itself a URL means a URL column; a name (often
        # hyperlinked to the profile) means a contact column.
        if samples and all(_URL_RE.match(str(s.get("text", ""))) for s in samples):
            return "contact_linkedin_url"
        return "contact"
    if "email" in words or "e mail" in h or "mail" in words:
        return "email"
    if words & {"phone", "telephone", "tel", "mobile", "number"}:
        return "phone"
    if words & {"website", "site", "domain", "url", "web"}:
        return "company_website"
    if words & {"company", "companies", "agency", "employer", "organisation", "organization", "business"}:
        return "company_name"
    if h in {"name", "contact", "contact name", "full name", "person", "recruiter"}:
        return "contact_name"
    if words & {"title", "position", "role", "job"}:
        return "contact_job_title"
    return "ignore"


def suggest_mapping(sheet: ParsedSheet) -> dict[str, Any]:
    columns = {letter: _role_for(sheet, letter, header) for letter, header in sheet.headers.items()}
    has_company = "company_name" in columns.values()
    has_people = bool({"contact", "contact_name", "phone", "email"} & set(columns.values()))
    return {
        "include": has_company or "company_website" in columns.values(),
        # Agency-style sheets list a company once, then its people underneath:
        # a blank company cell continues the group until a completely blank row.
        "group_fill": has_company and has_people,
        # On such sheets the company cell's hyperlink is the company's own site;
        # on a plain company list it is usually a job posting, not the employer.
        "company_link_is_website": has_company and has_people,
        "columns": columns,
    }
