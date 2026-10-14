"""Self-check for the import wizard's routes: `python test_contact_imports_routes.py`
(no pytest, no network).

The REAL auth gate (app.auth.get_current_user) runs on every request: tokens
are minted with a throwaway ES256 key served as the project's JWKS, two
admins come from ADMIN_USER_IDS, and a student is anyone else (the users
lookup is stubbed to "not an admin"). Supabase is an in-memory stand-in that
records what the router writes; the matching/commit SQL itself is tested in
supabase/tests/048_contact_import.test.sql and end to end in
test_contact_import_e2e.py.
"""

from __future__ import annotations

import io
import json
import os
import time
import uuid
from functools import lru_cache

SUPABASE_URL = "https://example.supabase.co"
os.environ["SUPABASE_URL"] = SUPABASE_URL
os.environ["SUPABASE_SERVICE_ROLE_KEY"] = "service-role-test-key"
os.environ["ADMIN_USER_IDS"] = "admin-a,admin-b"

import httpx  # noqa: E402
import openpyxl  # noqa: E402
from cryptography.hazmat.primitives import serialization  # noqa: E402
from cryptography.hazmat.primitives.asymmetric import ec  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from jose import jwk, jwt  # noqa: E402

from app import auth  # noqa: E402
from app.config import get_settings  # noqa: E402

get_settings.cache_clear()

from app.main import app  # noqa: E402
from app.routers import contact_imports  # noqa: E402

_failures: list[str] = []


def check(name: str, cond: bool, detail: object = "") -> None:
    print(f"[{'ok  ' if cond else 'FAIL'}] {name}{('  ' + repr(detail)) if detail != '' and not cond else ''}")
    if not cond:
        _failures.append(name)


# ── Auth: real verification against a throwaway key ──────────────────────────
_private = ec.generate_private_key(ec.SECP256R1())
_pem = _private.private_bytes(serialization.Encoding.PEM, serialization.PrivateFormat.PKCS8,
                              serialization.NoEncryption()).decode()
_pub = _private.public_key().public_bytes(serialization.Encoding.PEM,
                                         serialization.PublicFormat.SubjectPublicKeyInfo).decode()
_jwk = {k: (v.decode() if isinstance(v, bytes) else v) for k, v in jwk.construct(_pub, algorithm="ES256").to_dict().items()}


@lru_cache(maxsize=1)
def _fake_jwks(_url: str) -> dict:
    return {"k1": {**_jwk, "kid": "k1", "alg": "ES256"}}


auth._jwks = _fake_jwks


class _NotAdmin:
    status_code = 200

    def raise_for_status(self) -> None: ...

    def json(self) -> list:
        return [{"is_admin": False}]


auth.httpx.get = lambda *a, **k: _NotAdmin()   # the users lookup for anyone not in ADMIN_USER_IDS


def bearer(sub: str) -> dict[str, str]:
    token = jwt.encode({"sub": sub, "iss": f"{SUPABASE_URL}/auth/v1", "aud": "authenticated",
                        "exp": int(time.time()) + 3600}, _pem, algorithm="ES256", headers={"kid": "k1"})
    return {"Authorization": f"Bearer {token}"}


ADMIN_A, ADMIN_B, STUDENT = bearer("admin-a"), bearer("admin-b"), bearer("student-1")


# ── In-memory PostgREST stand-in ─────────────────────────────────────────────
def _match(row: dict, params: dict) -> bool:
    for key, cond in (params or {}).items():
        if key in ("select", "order", "limit", "offset", "or"):
            continue
        op, _, value = cond.partition(".")
        cell = row.get(key)
        if op == "eq" and str(cell) != value:
            return False
        if op == "in" and str(cell) not in value.strip("()").split(","):
            return False
    return True


class FakeRest:
    def __init__(self) -> None:
        self.tables: dict[str, list[dict]] = {"contact_import_batches": [], "contact_import_rows": [], "companies": []}
        self.rpcs: list[tuple[str, dict]] = []
        self.rpc_error: dict | None = None
        self._row_id = 0

    async def select(self, client, table, params=None, timeout=None):
        rows = [r for r in self.tables.get(table, []) if _match(r, params or {})]
        offset = int((params or {}).get("offset", 0))
        limit = int((params or {}).get("limit", 10**9))
        return [dict(r) for r in rows[offset:offset + limit]]

    async def count(self, client, table, params=None, timeout=None):
        return len([r for r in self.tables.get(table, []) if _match(r, params or {})])

    async def insert(self, client, table, rows, prefer="return=representation", timeout=None):
        out = []
        for row in rows:
            row = dict(row)
            if table == "contact_import_batches":
                row.setdefault("id", str(uuid.uuid4()))
                row.setdefault("status", "uploaded")
                row.setdefault("created_at", time.time())
            if table == "contact_import_rows":
                self._row_id += 1
                row.setdefault("id", self._row_id)
            self.tables.setdefault(table, []).append(row)
            out.append(dict(row))
        return out if prefer.startswith("return=representation") else []

    async def upsert(self, client, table, rows, on_conflict, prefer="", timeout=None):
        keys = on_conflict.split(",")
        for row in rows:
            existing = next((r for r in self.tables[table] if all(r.get(k) == row.get(k) for k in keys)), None)
            if existing is not None:
                existing.update(row)
            else:
                await self.insert(client, table, [row])
        return []

    async def update(self, client, table, match, values, prefer="return=minimal", timeout=None):
        hit = [r for r in self.tables.get(table, []) if _match(r, match)]
        for r in hit:
            r.update(values)
        return [dict(r) for r in hit] if prefer.startswith("return=representation") else []

    async def delete(self, client, table, match, prefer="return=representation", timeout=None):
        keep = [r for r in self.tables[table] if not _match(r, match)]
        gone = [r for r in self.tables[table] if _match(r, match)]
        self.tables[table] = keep
        return gone

    async def rpc(self, client, function, payload=None, timeout=None):
        self.rpcs.append((function, payload))
        if self.rpc_error is not None:
            status, body = self.rpc_error["status"], self.rpc_error["body"]
            request = httpx.Request("POST", f"{SUPABASE_URL}/rest/v1/rpc/{function}")
            raise httpx.HTTPStatusError("rpc failed", request=request,
                                        response=httpx.Response(status, json=body, request=request))
        if function == "contact_import_classify":
            for b in self.tables["contact_import_batches"]:
                if b["id"] == payload["p_batch"] and b["status"] == "uploaded":
                    b["status"] = "previewed"
            return {"total": 1, "matched": 0, "new": 1, "needs_review": 0, "invalid": 0, "skipped": 0}
        if function == "contact_import_commit":
            return {"status": "committed"}
        raise AssertionError(function)


fake = FakeRest()
contact_imports.SupabaseRest = lambda *a, **k: fake  # type: ignore[assignment]
client = TestClient(app)


def workbook() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "JOB AGENCIES UK"
    ws.append(["Company", "Link Job/Company", "LinkedIn", "Phone", "Email"])
    ws.append(["Acme Recruitment", None, "Jane Doe", "020 7946 0000", "jane@acme.example"])
    ws.append([None, None, "Tom Bell", None, "tom@acme.example"])
    ws.append([None, None, None, None, None])
    ws.append(["Beta Staffing", None, "Ann One", None, "ann@beta.example"])
    ws["C2"].hyperlink = "https://www.linkedin.com/in/jane-doe"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


XLSX = workbook()


def upload(headers: dict, data: bytes = XLSX, name: str = "contacts.xlsx"):
    return client.post("/contacts/imports", params={"filename": name}, content=data, headers=headers)


# ── 1. Access ────────────────────────────────────────────────────────────────
some_id = str(uuid.uuid4())
endpoints = [
    ("get", "/contacts/imports", None),
    ("post", "/contacts/imports?filename=x.xlsx", b"x"),
    ("get", "/contacts/imports/company-search?q=acme", None),
    ("get", f"/contacts/imports/{some_id}", None),
    ("get", f"/contacts/imports/{some_id}/rows", None),
    ("post", f"/contacts/imports/{some_id}/mapping", {"sheets": {}}),
    ("post", f"/contacts/imports/{some_id}/resolutions", {"resolutions": [{"row_id": 1, "clear": True}]}),
    ("post", f"/contacts/imports/{some_id}/commit", None),
]


def call(method: str, path: str, body, headers: dict | None):
    kwargs = {"headers": headers or {}}
    if isinstance(body, bytes):
        kwargs["content"] = body
    elif body is not None:
        kwargs["json"] = body
    return getattr(client, method)(path, **kwargs)


for method, path, body in endpoints:
    check(f"401 without a token: {method.upper()} {path.split('?')[0]}", call(method, path, body, None).status_code == 401)
    resp = call(method, path, body, STUDENT)
    check(f"403 for a signed-in non-admin: {method.upper()} {path.split('?')[0]}",
          resp.status_code == 403 and "coach access" in resp.json().get("detail", ""), resp.status_code)
check("every route was refused before touching the database", fake.rpcs == [] and fake.tables["contact_import_batches"] == [])

# ── 2. Upload ────────────────────────────────────────────────────────────────
resp = upload(ADMIN_A)
check("admin uploads a workbook → 201", resp.status_code == 201, resp.text[:300])
batch = resp.json()["batch"]
bid = batch["id"]
check("batch records the uploader, file kind and hash", batch["created_by"] == "admin-a" and batch["file_kind"] == "xlsx"
      and len(batch["file_sha256"]) == 64)
rows = [r for r in fake.tables["contact_import_rows"] if r["batch_id"] == bid]
check("only non-blank rows are stored, raw cells intact",
      [r["row_number"] for r in rows] == [2, 3, 5] and rows[0]["raw"]["C"]["hyperlink"] == "https://www.linkedin.com/in/jane-doe")
check("the sheet summary carries headers and a suggested mapping",
      batch["sheets"][0]["suggested_mapping"]["columns"]["C"] == "contact" and batch["sheets"][0]["suggested_mapping"]["group_fill"])
check("first upload of this file: no previous imports", resp.json()["previous_imports_of_this_file"] == [])
again = upload(ADMIN_A).json()
check("uploading the same file again reports the earlier import",
      [p["id"] for p in again["previous_imports_of_this_file"]] == [bid])

check("wrong file type → 422", upload(ADMIN_A, b"hello", "notes.txt").status_code == 422)
check("unreadable workbook → 422", upload(ADMIN_A, b"PK nope", "c.xlsx").status_code == 422)
big = client.post("/contacts/imports", params={"filename": "big.csv"}, content=b"x" * (10 * 1024 * 1024 + 1), headers=ADMIN_A)
check("over 10 MB → 413", big.status_code == 413, big.status_code)
check("header-only file → 422", upload(ADMIN_A, b"Company,Email\n", "empty.csv").status_code == 422)

# ── 3. Cross-admin: history is shared, changes are the uploader's ────────────
resp = client.get(f"/contacts/imports/{bid}", headers=ADMIN_B)
check("another admin can read the batch (mine = false)", resp.status_code == 200 and resp.json()["batch"]["mine"] is False)
check("…and sees it in the history", any(b["id"] == bid and not b["mine"] for b in client.get("/contacts/imports", headers=ADMIN_B).json()))
mapping = {"sheets": {"JOB AGENCIES UK": batch["sheets"][0]["suggested_mapping"]}}
for label, method, path, body in [
    ("mapping", "post", f"/contacts/imports/{bid}/mapping", mapping),
    ("resolutions", "post", f"/contacts/imports/{bid}/resolutions", {"resolutions": [{"row_id": rows[0]["id"], "clear": True}]}),
    ("commit", "post", f"/contacts/imports/{bid}/commit", None),
]:
    resp = call(method, path, body, ADMIN_B)
    check(f"another admin cannot {label} this batch → 403", resp.status_code == 403, resp.status_code)
check("unknown batch → 404", client.get(f"/contacts/imports/{uuid.uuid4()}", headers=ADMIN_A).status_code == 404)
check("malformed batch id → 404, not a database error", client.get("/contacts/imports/not-a-uuid", headers=ADMIN_A).status_code == 404)

# ── 4. Mapping → preview ─────────────────────────────────────────────────────
resp = client.post(f"/contacts/imports/{bid}/resolutions", json={"resolutions": [{"row_id": rows[0]["id"], "clear": True}]}, headers=ADMIN_A)
check("decisions before a mapping → 409", resp.status_code == 409)
bad = {"sheets": {"JOB AGENCIES UK": {**mapping["sheets"]["JOB AGENCIES UK"], "columns": {"A": "company_name", "B": "salary"}}}}
check("unknown role → 422", client.post(f"/contacts/imports/{bid}/mapping", json=bad, headers=ADMIN_A).status_code == 422)
resp = client.post(f"/contacts/imports/{bid}/mapping", json=mapping, headers=ADMIN_A)
check("mapping applied → preview returned", resp.status_code == 200 and resp.json()["preview"]["new"] == 1, resp.text[:200])
stored = {r["row_number"]: r for r in fake.tables["contact_import_rows"] if r["batch_id"] == bid}
check("extraction written for every row; grouping applied",
      stored[3]["extracted"]["company"]["name"] == "Acme Recruitment" and stored[5]["extracted"]["company"]["name"] == "Beta Staffing")
check("a new mapping clears earlier decisions", all(r["resolution"] is None for r in stored.values()))
check("the preview ran in the database", fake.rpcs[-1] == ("contact_import_classify", {"p_batch": bid}))

# ── 5. Resolutions ───────────────────────────────────────────────────────────
row_id = stored[2]["id"]
resp = client.post(f"/contacts/imports/{bid}/resolutions",
                   json={"resolutions": [{"row_id": row_id, "company": {"action": "match"}}]}, headers=ADMIN_A)
check("match without a company → 422", resp.status_code == 422)
resp = client.post(f"/contacts/imports/{bid}/resolutions", json={"resolutions": [{"row_id": 99999, "clear": True}]}, headers=ADMIN_A)
check("a row from another batch → 404", resp.status_code == 404)
cid = str(uuid.uuid4())
resp = client.post(f"/contacts/imports/{bid}/resolutions", json={"resolutions": [
    {"row_id": row_id, "company": {"action": "match", "company_id": cid, "save_alias": True}, "contact": {"action": "skip"}}]},
    headers=ADMIN_A)
check("decision stored as the database expects it",
      resp.status_code == 200 and stored[2]["resolution"] == {
          "company": {"action": "match", "company_id": cid, "save_alias": True}, "contact": {"action": "skip"}}, stored[2]["resolution"])

# ── 6. Commit and its error mapping ──────────────────────────────────────────
for code, status in [("CI409", 409), ("CI403", 403), ("CI404", 404), ("CI400", 409), ("XX000", 502)]:
    fake.rpc_error = {"status": 400, "body": {"code": code, "message": "boom", "details": None, "hint": None}}
    resp = client.post(f"/contacts/imports/{bid}/commit", headers=ADMIN_A)
    check(f"commit: database {code} → HTTP {status}", resp.status_code == status, resp.status_code)
fake.rpc_error = None
resp = client.post(f"/contacts/imports/{bid}/commit", headers=ADMIN_A)
check("commit passes the uploader as the actor", resp.status_code == 200
      and fake.rpcs[-1] == ("contact_import_commit", {"p_batch": bid, "p_actor": "admin-a"}))
fake.tables["contact_import_batches"][0]["status"] = "committed"
check("a committed batch can't be re-mapped → 409",
      client.post(f"/contacts/imports/{bid}/mapping", json=mapping, headers=ADMIN_A).status_code == 409)

# ── 7. Reads and routing ─────────────────────────────────────────────────────
resp = client.get(f"/contacts/imports/{bid}/rows", params={"classification": "bogus"}, headers=ADMIN_A)
check("unknown classification filter → 422", resp.status_code == 422)
resp = client.get(f"/contacts/imports/{bid}/rows", params={"limit": 2}, headers=ADMIN_A)
check("rows are paged with a total", resp.status_code == 200 and len(resp.json()["rows"]) == 2 and resp.json()["total"] == 3)
fake.tables["companies"].append({"id": cid, "name": "Acme Recruitment", "slug": "acme-recruitment"})
resp = client.get("/contacts/imports/company-search", params={"q": "ac(me),*"}, headers=ADMIN_A)
check("company search survives PostgREST metacharacters", resp.status_code == 200)
resp = client.post("/contacts", json={}, headers=ADMIN_A)
check("the contacts router still owns POST /contacts (422 from its own validation)", resp.status_code == 422, resp.status_code)

print()
if _failures:
    print(f"{len(_failures)} FAILED")
    raise SystemExit(1)
print("ALL CONTACT IMPORT ROUTE CHECKS PASSED")
